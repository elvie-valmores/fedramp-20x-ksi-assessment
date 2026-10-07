#!/usr/bin/env python3
"""Emits the Security Decision Record: the project's actual deliverable.

Reads the 46 determinations from docs/KSI-Design-Matrix.xlsx and the
collector's latest results, and writes one SDR in FedRAMP's schema, as JSON
and as a human-readable rendering of that same JSON. Validates against the
pinned schema before writing anything; a validation error fails the run.

    cd sdr && ../.venv/bin/python emit.py --results ../collector/results.json

See docs/SDR-Emitter-Spec.md for the mapping, and DECISIONS.md, 2026-09-23,
for the choices made where the spec left one open.

What the output does not claim, stated here and in the rendering:

    No independent assessor is engaged. ksiAssessment carries the provider's
    own assurance reasoning and says so, rather than wearing an assessor's
    label.

    fedRampRequirements is emitted empty. The project determined the 46
    indicators and not the FRR rules; --frr is a switch so that the omission
    is a stated choice, not a silent one.

    SDR-CSX-KMT's historical metrics come from the record store: since
    2026-10-01 each daily run in GitHub Actions (collect.yml) writes its
    results to the log store under collector-runs/, and --history hands
    them to this emitter. Without --history a record carries one run and
    says so. Collection began 2026-10-01, so "up to the past year" is what
    exists. --runtime says whether this run
    came from that schedule or from a laptop, so the cycle statement is never
    wrong about which. The collection window is stated on every indicator.

    certificationPackageOverviewUri points at the repository README at the
    emitting commit. It is the project's overview. It is not a FedRAMP
    Certification Package Overview, and the rendering says so.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
DOCS = REPO / "docs"
MATRIX = DOCS / "KSI-Design-Matrix.xlsx"
OUT = Path(__file__).resolve().parent / "out"
REPOSITORY_URL = "https://github.com/elvie-valmores/fedramp-20x-ksi-assessment"

# Pinned by filename, internal version and hash. The filename is dated and
# does not change when the content does -- the SDR schema moved from 1.0.3
# to 1.1.1 under the same name -- so all three are checked.
SCHEMAS = {
    "sdr": ("fedramp-security-decision-record-schema-2026-06-24.json", "1.1.1",
            "94580404e8d4276ee3ced2a3c39cce5cc6565df6ae38a91f60f26cb96d833188"),
    "common": ("fedramp-common-definitions-schema-2026-06-24.json", "0.4.0",
               "be90c62c8dd8d270d48965b92e3a2a831b8b81791d509cc92557b72ccb0a1b44"),
}

# A project-level fact, not a per-row one: how the run behind this record
# was produced. Chosen by --runtime, which has no default.
CYCLE_PRACTICE = {
    "ci": (
        "Current practice: the collector runs daily from GitHub Actions "
        "(.github/workflows/collect.yml), as a read-only identity in each cloud "
        "with no stored credential, and this record was emitted by that run. "
        "Daily is stricter than every required cycle above (DECISIONS.md, 2026-10-01)."
    ),
    "local": (
        "Current practice: this record came from a collector run on an operator's "
        "workstation, not from the daily scheduled run in GitHub Actions "
        "(.github/workflows/collect.yml), which is the system's cycle "
        "(DECISIONS.md, 2026-10-01)."
    ),
}
NO_ASSESSOR = (
    "**No independent assessor is engaged.** The statements below are the "
    "provider's own assurance reasoning from the determination. They are not an "
    "independent assessment and should not be read as one."
)
STATUS = {
    "IN SCOPE": "Implemented",
    "IN SCOPE, PARTIALLY SATISFIED": "Partially Implemented",
    "DEFERRED": "Not Implemented",
}


# --- reading the determinations ---


@dataclass
class Determination:
    ksi_id: str
    title: str
    status: str
    sections: dict[str, str] = field(default_factory=dict)
    tables: dict[str, list[list[str]]] = field(default_factory=dict)

    @property
    def deferred(self) -> bool:
        return self.status == "DEFERRED"


TITLE = re.compile(r"^(KSI-[A-Z]{3}-[A-Z]{3}) — (.+?)\s{2,}(" + "|".join(map(re.escape, STATUS)) + r")\s*$")
LABEL = re.compile(
    r"^(Requirement(?: \(Class C variant\))?|Design rationale(?: for deferral)?|"
    r"Automation assurance|Remediation|Limitations and residual risk):\s+(.*)$",
    re.S,
)
TABLE = re.compile(r"^(BUILD|VERIFY|VALIDATE|PROVE) — ")
NOT_DETERMINATIONS = {"Index", "Responsibility Matrix", "Assertion Coverage"}


def read_determinations() -> list[Determination]:
    """Every determination in the matrix, or an error naming what was not understood.

    Anything inside a determination block that is not a known label, table
    header or table row stops the run. A parser that skipped what it did not
    recognise would drop text from the SDR without anyone knowing.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(MATRIX, read_only=True)
    found, problems = [], []
    for sheet in workbook.worksheets:
        if sheet.title in NOT_DETERMINATIONS:
            continue
        current, table = None, None
        for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = ["" if c is None else str(c).strip() for c in row]
            first = cells[0] if cells else ""
            if not any(cells):
                continue
            if title := TITLE.match(first):
                current = Determination(title.group(1), title.group(2).strip(), title.group(3))
                found.append(current)
                table = None
            elif current is None:
                continue  # the sheet's own heading and introduction
            elif label := LABEL.match(first):
                key = "Requirement" if label.group(1).startswith("Requirement") else label.group(1)
                current.sections[key] = label.group(2).strip()
                table = None
            elif header := TABLE.match(first):
                table = header.group(1)
                current.tables.setdefault(table, [])
            elif table:
                current.tables[table].append(cells[:5])
            else:
                problems.append(f"{sheet.title}!R{number}: not understood: {first[:60]!r}")
    if problems:
        sys.exit("the matrix has content the emitter does not understand:\n  " + "\n  ".join(problems))
    return found


# --- building the record ---


def build_statement(row: list[str]) -> str:
    number, what, configuration, cloud, provisioned = row
    return f"**{number}. {what}** — {configuration} *({cloud}; provisioned by {provisioned})*"


def evidence_row_statement(row: list[str], row_id: str | None = None, coverage: dict | None = None) -> str:
    kind, artifact, source, required, implemented = row
    statement = f"**{kind}** — {artifact}. Source: {source}. Required every {required}; implemented: {implemented}."
    if row_id is None or coverage is None:
        return statement
    # Which checks prove this row, from the links on the checks themselves
    # (sdr/matrix_rows.py). A row with no check says so rather than leaving
    # the reader to infer it from the evidence list.
    entry = coverage[row_id]
    checks = ", ".join(
        f"`{c}`" + (f" (judged only while {entry['environments'][c]} stands)" if c in entry.get("environments", {}) else "")
        for c in entry["checks"])
    if entry["status"] == "full":
        return f"{statement} *Row `{row_id}`: automated by {checks}.*"
    if entry["status"] == "partial":
        # Two checks with the same gap state it once.
        gaps = " ".join(dict.fromkeys(g.split(": ", 1)[1] for g in entry["gaps"]))
        return f"{statement} *Row `{row_id}`: partly automated by {checks}. Not covered: {gaps}*"
    if entry["status"] == "recorded":
        p = entry["position"]
        where = f" Evidence: {p['evidence']}." if p.get("evidence") else ""
        return (f"{statement} *Row `{row_id}`: not automated -- {p['position']}. {p['reason']}{where} "
                f"Risk: {p['risk']}*")
    return f"{statement} *Row `{row_id}`: no automated check yet.*"


def indicator(det: Determination, outcomes: list[dict], context: dict, controls: dict) -> dict:
    mine = [o for o in outcomes if o["check"]["indicator"] == det.ksi_id]
    s = det.sections

    if det.deferred:
        implementation = [
            f"**Deferred.** {s['Design rationale for deferral']}",
            f"**Resulting risk.** {s['Limitations and residual risk']}",
        ]
    else:
        implementation = [s["Design rationale"]] + [build_statement(r) for r in det.tables.get("BUILD", [])]

    cadences = sorted({r[3] for t in ("VERIFY", "VALIDATE", "PROVE") for r in det.tables.get(t, []) if r[3]})
    validation = [
        f"**Cycle.** Required: {', '.join(cadences) or 'not stated'}. {CYCLE_PRACTICE[context['runtime']]}",
        *[evidence_row_statement(r, f"{det.ksi_id}.{t.lower()}.{n}", context["coverage"])
          for t in ("VERIFY", "VALIDATE") for n, r in enumerate(det.tables.get(t, []), start=1)],
        *[evidence_row_statement(r) for r in det.tables.get("PROVE", [])],
        collection_statement(mine, context),
    ]

    assessment = [NO_ASSESSOR]
    if "Automation assurance" in s:
        assessment.append(f"**Automation assurance.** {s['Automation assurance']}")
    assessment.append(f"**Limitations and residual risk.** {s['Limitations and residual risk']}")

    return {
        "ksiId": det.ksi_id,
        "ksiImplementationStatus": STATUS[det.status],
        "ksiImplementation": implementation,
        "ksiValidation": validation,
        "ksiAssessment": assessment,
        "ksiTests": tests(det, mine, controls),
        "ksiEvidence": [evidence(o, context) for o in mine]
        + [m for m in [metrics_evidence(det.ksi_id, context["history"], date.fromisoformat(context["run_date"]))]
           if m and context["history"]],
    }


def collection_statement(mine: list[dict], context: dict) -> str:
    if not mine:
        return (
            "**Automated evidence.** No collector check exists for this indicator yet, so no "
            "evidence is attached. The rows above are the determination's design."
        )
    counts = {}
    for o in mine:
        counts[o["status"]] = counts.get(o["status"], 0) + 1
    tally = ", ".join(f"{n} {k}" for k, n in sorted(counts.items()))
    if context["history"]:
        first = min(context["history"])
        return (
            f"**Automated evidence.** {len(mine)} collector check(s) in the run on "
            f"{context['run_date']}: {tally}. Historical metrics (SDR-CSX-KMT) cover "
            f"{len(context['history'])} day(s) of CI runs since {first}, in the metrics "
            "evidence object below; the year is what exists, not a full year."
        )
    return (
        f"**Automated evidence.** {len(mine)} collector check(s) in one run on "
        f"{context['run_date']}: {tally}. Collection window: that single run. This record was "
        "emitted without --history, so SDR-CSX-KMT's 30-day and yearly summaries are not included."
    )


# --- historical metrics (SDR-CSX-KMT) ---
#
# Class C asks, per indicator, for a 30-day summary, a summary up to the past
# year, and all daily data up to the past year, each "where available". The
# data is the record store: every CI run writes its results to the
# log store under collector-runs/ (Object Locked, evidence key), and
# collect.yml hands those records to this emitter with --history. One data
# point per day: the latest CI run that day, scheduled or started by hand. Workstation runs are never
# stored or counted. See DECISIONS.md, 2026-10-01.


def load_history(directory: Path) -> dict[date, dict]:
    """{day: that day's latest CI run record} from a directory of records."""
    latest: dict[date, dict] = {}
    for path in sorted(directory.rglob("*.json")):
        run = json.loads(path.read_text())
        if (run.get("run") or {}).get("runtime") != "ci":
            continue
        day = date.fromisoformat(run["started_at"][:10])
        if day not in latest or run["started_at"] > latest[day]["started_at"]:
            latest[day] = run
    return latest


def _day_counts(run: dict, ksi_id: str) -> dict | None:
    everything = [o for o in run["outcomes"] if o["check"]["indicator"] == ksi_id]
    # A check whose environment was torn down was not judged that day. It is
    # left out of the day's count rather than counted as an error: the
    # environment is down by design between sessions (collector/environments.py).
    mine = [o for o in everything if o["status"] != "NOT_STANDING"]
    if not mine:
        return None
    return {
        "checks": len(mine),
        "not_standing": len(everything) - len(mine),
        "passed": sum(o["status"] == "PASS" for o in mine),
        "failed": sum(o["status"] == "FAIL" for o in mine),
        "errored": sum(o["status"] not in ("PASS", "FAIL") for o in mine),
        "not_passing": sorted(o["check"]["id"] for o in mine if o["status"] != "PASS"),
    }


def _window(days: dict[date, dict], as_of: date, length: int) -> str:
    start = as_of - timedelta(days=length - 1)
    within = {d: c for d, c in days.items() if start <= d <= as_of}
    if not within:
        return f"no data in the {length} days to {as_of}"
    clean = sum(c["passed"] == c["checks"] for c in within.values())
    rate = sum(c["passed"] / c["checks"] for c in within.values()) / len(within)
    misses: dict[str, int] = {}
    for c in within.values():
        for check in c["not_passing"]:
            misses[check] = misses.get(check, 0) + 1
    worst = ", ".join(f"{k} on {n} day(s)" for k, n in sorted(misses.items(), key=lambda x: (-x[1], x[0])))
    return (
        f"data on {len(within)} of {length} days ({min(within)} to {max(within)}); every check passed on "
        f"{clean} of {len(within)}; mean daily pass rate {rate:.0%}"
        + (f"; not passing: {worst}" if worst else "")
    )


def metrics_evidence(ksi_id: str, history: dict[date, dict], as_of: date) -> dict | None:
    """The SDR-CSX-KMT evidence object for one indicator, or None without data."""
    days = {d: c for d, run in history.items() if (c := _day_counts(run, ksi_id))}
    if not days:
        return None
    year = {d: c for d, c in days.items() if as_of - timedelta(days=364) <= d <= as_of}
    daily = "\n".join(
        f"{d}: {c['passed']}/{c['checks']} passed"
        + (f" (not passing: {', '.join(c['not_passing'])})" if c["not_passing"] else "")
        for d, c in sorted(year.items())
    )
    return {
        # The schema's evidenceType is an enumeration; a metrics summary is a
        # Report, and the description says which kind.
        "evidenceType": "Report",
        "evidenceDescription": (
            f"Historical metrics (SDR-CSX-KMT): daily collector results from CI for {ksi_id} (the daily schedule and any run started by hand in CI; never a workstation run). Past 30 days: {_window(days, as_of, 30)}. "
            f"Past year: {_window(days, as_of, 365)}. The record store began {min(days)}, so the year is "
            "what exists, not a full year."
        ),
        "evidenceText": "All daily data, past year (the latest CI run each day):\n" + daily,
        "lastUpdated": max(days).isoformat(),
    }


def tests(det: Determination, mine: list[dict], controls: dict) -> list[str]:
    found = []
    for o in mine:
        check = o["check"]
        key = check["params"].get("assertion") or check["params"].get("resource") or check["params"].get("judge")
        if (check["mechanism"], key) in controls:
            found.append(
                f"Negative control for `{check['id']}` (collector/self_test.py): "
                f"{controls[(check['mechanism'], key)]}."
            )
    # Test scenarios the determination designs but nothing runs yet. Listed
    # as designed, because a test that has not run has not validated anything.
    for row in det.tables.get("VALIDATE", []):
        if re.search(r"deliberate|test", f"{row[1]} {row[2]}", re.I):
            found.append(f"Designed, not yet automated: {row[1]} ({row[2]}).")
    if not found:
        found.append("None run or designed." if det.deferred else "No test is automated for this indicator yet.")
    return found


def evidence(outcome: dict, context: dict) -> dict:
    check, result = outcome["check"], outcome.get("result") or {}
    if check["evidence_type"] == "CFG":
        kind = "Configuration"
    elif check["mechanism"] == "log_query":
        kind = "Log"
    else:
        kind = "Report"
    ran = (result.get("ran_at") or context["started_at"])[:10]
    return {
        "evidenceType": kind,
        "evidenceDescription": f"{check['description']}. Result: {outcome['status']} — {outcome['message']}",
        "evidenceLocation": f"{REPOSITORY_URL}/blob/{context['commit']}/collector/checks/{check['id']}.json",
        "evidenceText": json.dumps(
            {"check": check["id"], "status": outcome["status"], "evidence": result.get("evidence")},
            indent=2, default=str,
        ),
        "lastUpdated": ran,
    }


def record(dets, outcomes, context, controls, frr: str) -> dict:
    if frr != "empty":
        raise NotImplementedError(
            "--frr populated needs FRR determinations (MAS, VDR, IVV, SDR), which do not exist yet"
        )
    return {
        "certificationPackageOverviewUri": f"{REPOSITORY_URL}/blob/{context['commit']}/README.md",
        "metadata": {
            "version": context["version"],
            "lastUpdated": context["emitted_at"],
            "updateSource": (
                "Automated: sdr/emit.py, from docs/KSI-Design-Matrix.xlsx and the collector run "
                f"started {context['started_at']}"
            ),
        },
        "fedRampRequirements": [],
        "keySecurityIndicators": [indicator(d, outcomes, context, controls) for d in dets],
    }


# --- checking what was built ---


def load_schemas() -> tuple[dict, dict]:
    loaded = {}
    for key, (name, version, digest) in SCHEMAS.items():
        raw = (DOCS / name).read_bytes()
        actual = hashlib.sha256(raw).hexdigest()
        schema = json.loads(raw)
        if actual != digest or schema.get("$schemaVersion") != version:
            sys.exit(
                f"{name} does not match its pin: sha256 {actual[:12]}…, "
                f"$schemaVersion {schema.get('$schemaVersion')} (pinned {digest[:12]}…, {version})"
            )
        loaded[key] = schema
    return loaded["sdr"], loaded["common"]


def validate(sdr: dict, schema: dict, common: dict) -> list[str]:
    from jsonschema import Draft202012Validator, FormatChecker
    from referencing import Registry, Resource

    registry = Registry().with_resource(common["$id"], Resource.from_contents(common))
    validator = Draft202012Validator(schema, registry=registry, format_checker=FormatChecker())
    problems = [f"{'/'.join(map(str, e.absolute_path))}: {e.message}" for e in validator.iter_errors(sdr)]

    # jsonschema checks "uri" and "date-time" formats only when optional
    # packages are installed, so a pass would not mean they were checked.
    # These are checked here instead, and so is the spec's own rule that an
    # evidence object is never empty even though the schema allows it.
    uri = re.compile(r"^https://\S+$")
    if not uri.match(sdr["certificationPackageOverviewUri"]):
        problems.append("certificationPackageOverviewUri is not an https URI")
    datetime.fromisoformat(sdr["metadata"]["lastUpdated"])
    for k in sdr["keySecurityIndicators"]:
        for i, e in enumerate(k["ksiEvidence"]):
            where = f"{k['ksiId']} evidence {i}"
            if not (e.get("evidenceType") and e.get("evidenceDescription")):
                problems.append(f"{where}: missing type or description")
            if not (e.get("evidenceLocation") or e.get("evidenceText")):
                problems.append(f"{where}: neither location nor text")
            if e.get("evidenceLocation") and not uri.match(e["evidenceLocation"]):
                problems.append(f"{where}: evidenceLocation is not an https URI")
            date.fromisoformat(e["lastUpdated"])
    return problems


# --- the rendering ---


def render(sdr: dict, context: dict) -> str:
    """Markdown generated from the emitted JSON, never from the matrix.

    Rendering the record rather than re-deriving it means the two formats
    cannot say different things.
    """
    ksis = sdr["keySecurityIndicators"]
    counts = {}
    for k in ksis:
        counts[k["ksiImplementationStatus"]] = counts.get(k["ksiImplementationStatus"], 0) + 1
    evidence_total = sum(len(k["ksiEvidence"]) for k in ksis)
    covered = sum(1 for k in ksis if k["ksiEvidence"])

    lines = [
        "# Security Decision Record — Caliper",
        "",
        f"Version `{sdr['metadata']['version']}`, emitted {sdr['metadata']['lastUpdated']}. "
        f"Schema {SCHEMAS['sdr'][0]} `$schemaVersion` {SCHEMAS['sdr'][1]}; validated on emit.",
        "",
        "## What this record does and does not claim",
        "",
        f"- **{len(ksis)} indicators**: " + ", ".join(f"{n} {s}" for s, n in sorted(counts.items())) + ".",
        f"- **Automated evidence** for {covered} of {len(ksis)} indicators, {evidence_total} evidence objects, "
        f"from one collector run started {context['started_at']}.",
        "- **Evidence rows:** " + ", ".join(
            f"{sum(1 for v in context['coverage'].values() if v['status'] == s)} {label}"
            for s, label in (("full", "automated"), ("partial", "partly automated"),
                             ("recorded", "recorded as descoped, manual or a stated gap"), ("none", "not yet automated"))
        ) + f", of {len(context['coverage'])}. Each row in Validation names the checks behind it.",
        "- **No independent assessor is engaged.** Assessment statements are the provider's own reasoning.",
        "- **`fedRampRequirements` is empty.** The project determined the 46 indicators, not the FRR rules.",
        (f"- **Historical metrics (SDR-CSX-KMT)** from {len(context['history'])} day(s) of CI runs "
         f"since {min(context['history'])}, per indicator in a metrics evidence object. The year is what "
         "exists since then, not a full year."
         if context["history"] else
         "- **No historical metrics (SDR-CSX-KMT) in this record.** It was emitted without --history, "
         "so the collection window is one run."),
        f"- **`certificationPackageOverviewUri`** points at the project README. It is not a FedRAMP "
        "Certification Package Overview.",
        "",
    ]
    for k in ksis:
        lines += [f"## {k['ksiId']} — {k['ksiImplementationStatus']}", ""]
        for heading, key in (("Implementation", "ksiImplementation"), ("Validation", "ksiValidation"),
                             ("Assessment", "ksiAssessment")):
            lines += [f"### {heading}", ""] + [f"- {x}" for x in k[key]] + [""]
        lines += ["### Tests", ""] + [f"- {x}" for x in k["ksiTests"]] + [""]
        lines += ["### Evidence", ""]
        if not k["ksiEvidence"]:
            lines += ["None attached.", ""]
        for e in k["ksiEvidence"]:
            link = f" ([definition]({e['evidenceLocation']}))" if e.get("evidenceLocation") else ""
            lines += [f"- *{e['evidenceType']}, {e['lastUpdated']}* — {e['evidenceDescription']}{link}"]
        lines += [""]
    return "\n".join(lines)


# --- running it ---


def git(*args: str) -> str:
    return subprocess.run(["git", *args], cwd=REPO, capture_output=True, text=True, check=True).stdout.strip()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--results", type=Path, required=True, help="collector output from run_checks.py --json")
    parser.add_argument("--frr", choices=["empty", "populated"], required=True,
                        help="how to emit fedRampRequirements; only 'empty' is possible today")
    parser.add_argument("--runtime", choices=["ci", "local"], required=True,
                        help="where the collector run came from: the scheduled CI job, or a workstation")
    parser.add_argument("--history", type=Path,
                        help="directory of stored CI run records, for SDR-CSX-KMT's metrics")
    args = parser.parse_args()

    schema, common = load_schemas()
    run = json.loads(args.results.read_text())
    commit = git("rev-parse", "HEAD")
    changed = git("status", "--porcelain")
    dirty = bool(changed)
    if dirty:
        # Name what made it dirty, so a "-dirty" version is never a mystery.
        print("checkout has uncommitted changes; the version is marked -dirty:\n" + changed, file=sys.stderr)
    context = {
        "commit": commit,
        # Evidence links point at this commit. If the tree has uncommitted
        # changes, the links may not show what was run, and the version says so.
        "version": commit[:12] + ("-dirty" if dirty else ""),
        "emitted_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "started_at": run["started_at"],
        "run_date": run["started_at"][:10],
        "runtime": args.runtime,
        "history": load_history(args.history) if args.history else {},
    }

    sys.path.insert(0, str(REPO / "collector"))
    import self_test

    dets = read_determinations()
    ids = [d.ksi_id for d in dets]
    if len(ids) != len(set(ids)):
        sys.exit("a KSI appears twice in the matrix")
    unknown = sorted({o["check"]["indicator"] for o in run["outcomes"]} - set(ids))
    if unknown:
        # Evidence for an indicator the SDR does not contain would vanish.
        sys.exit(f"collector results name indicators the matrix does not: {', '.join(unknown)}")

    from matrix_rows import coverage, rows_from

    context["coverage"] = coverage([o["check"] for o in run["outcomes"]], rows_from(dets))
    sdr = record(dets, run["outcomes"], context, self_test.negative_controls(), args.frr)
    problems = validate(sdr, schema, common)
    if problems:
        print("the record does not validate; nothing written:", *problems[:40], sep="\n  ")
        return 1

    OUT.mkdir(exist_ok=True)
    (OUT / "sdr.json").write_text(json.dumps(sdr, indent=2) + "\n")
    (OUT / "sdr.md").write_text(render(sdr, context))
    print(f"wrote sdr/out/sdr.json and sdr/out/sdr.md: {len(dets)} indicators, "
          f"{sum(len(k['ksiEvidence']) for k in sdr['keySecurityIndicators'])} evidence objects, "
          f"version {context['version']}, valid against {SCHEMAS['sdr'][0]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
